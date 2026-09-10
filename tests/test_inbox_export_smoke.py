"""Smoke: POST /api/inbox/export + ai_wrong filter + ai_wrong_note serialize."""
from __future__ import annotations

from datetime import datetime, timezone
from uuid import uuid4

import bcrypt
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session, sessionmaker

from app.api.inbox_export import AI_REVIEW_PRESET
from app.api.main import app
from app.db.models import Document, Lot, LotState, Run, ScoutSession, User
from tests.conftest import SMOKE_PREFIX

_PASS = "qa-smoke-export-pass"


def _hash(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def _client() -> TestClient:
    try:
        return TestClient(app, lifespan="off")
    except TypeError:
        return TestClient(app)


def _cleanup(
    factory: sessionmaker[Session],
    *,
    username: str,
    lot_ids: list[str],
    query: str,
) -> None:
    with factory() as session:
        user = session.scalar(select(User).where(User.username == username))
        if user is not None:
            session.execute(delete(ScoutSession).where(ScoutSession.user_id == user.id))
            session.delete(user)
        if lot_ids:
            session.execute(delete(Document).where(Document.tender_id.in_(lot_ids)))
            session.execute(delete(LotState).where(LotState.tender_id.in_(lot_ids)))
            session.execute(delete(Lot).where(Lot.tender_id.in_(lot_ids)))
        session.execute(delete(Run).where(Run.query == query))
        session.commit()


@pytest.mark.smoke
def test_inbox_export_csv_xlsx_ai_wrong(smoke_db: sessionmaker[Session]) -> None:
    suffix = uuid4().hex[:12]
    username = f"{SMOKE_PREFIX}export_{suffix}"
    visible_id = f"{SMOKE_PREFIX}exp_vis_{suffix}"
    wrong_id = f"{SMOKE_PREFIX}exp_wrong_{suffix}"
    hidden_id = f"{SMOKE_PREFIX}exp_hid_{suffix}"
    query = f"{SMOKE_PREFIX}export_run_{suffix}"
    lot_ids = [visible_id, wrong_id, hidden_id]
    try:
        with smoke_db() as session:
            session.add(
                User(
                    username=username,
                    password_hash=_hash(_PASS),
                    display_name="qa_smoke_export",
                )
            )
            session.add(Run(query=query, status="done", limit_n=10))
            session.flush()
            run = session.scalar(select(Run).where(Run.query == query))
            assert run is not None
            for tid, title in (
                (visible_id, "УЗК экспорт visible"),
                (wrong_id, "УЗК экспорт ai wrong"),
                (hidden_id, "УЗК экспорт hidden"),
            ):
                session.add(
                    Lot(
                        tender_id=tid,
                        run_id=run.id,
                        title=title,
                        url=f"https://rostender.info/tender/{tid}",
                        score=6,
                        tier="L1",
                        location="Казань",
                        customer_name="ООО Export",
                        deadline_msk="20.12.2026",
                        status="Приём заявок",
                        fit_reason="услуга НК",
                        source_platform_id="rostender",
                        ingested_at=datetime(2026, 9, 1, 12, 0, tzinfo=timezone.utc),
                    )
                )
            session.flush()
            session.add(
                LotState(
                    tender_id=wrong_id,
                    ai_reviewed_at=datetime(2026, 9, 2, 10, 0, tzinfo=timezone.utc),
                    ai_tier="L2",
                    ai_reason_ru="тест",
                    ai_wrong_at=datetime(2026, 9, 3, 10, 0, tzinfo=timezone.utc),
                    ai_wrong_note="модель ошиблась",
                    ai_trigger="manual",
                )
            )
            session.add(
                LotState(
                    tender_id=hidden_id,
                    board_hidden=True,
                )
            )
            session.add(
                Document(
                    tender_id=visible_id,
                    filename="spec.pdf",
                    stored_path="noop",
                    size_bytes=100,
                )
            )
            session.commit()

        with _client() as client:
            login = client.post(
                "/api/auth/login",
                json={"username": username, "password": _PASS},
            )
            assert login.status_code == 200

            listing = client.get("/api/inbox")
            assert listing.status_code == 200
            ids = [row["tender_id"] for row in listing.json()["items"]]
            assert visible_id in ids
            assert wrong_id in ids
            assert hidden_id not in ids
            wrong_row = next(r for r in listing.json()["items"] if r["tender_id"] == wrong_id)
            assert wrong_row["ai_wrong"] is True
            assert wrong_row["ai_wrong_note"] == "модель ошиблась"

            only_wrong = client.get("/api/inbox", params={"ai_wrong": "1"})
            assert only_wrong.status_code == 200
            wrong_ids = [row["tender_id"] for row in only_wrong.json()["items"]]
            assert wrong_id in wrong_ids
            assert visible_id not in wrong_ids

            bad = client.post(
                "/api/inbox/export",
                json={"format": "xlsm", "columns": list(AI_REVIEW_PRESET)},
            )
            assert bad.status_code == 400

            csv_res = client.post(
                "/api/inbox/export",
                json={
                    "format": "csv",
                    "columns": ["tender_id", "title", "ai_wrong", "ai_wrong_note", "documents_count"],
                },
            )
            assert csv_res.status_code == 200
            assert csv_res.content.startswith(b"\xef\xbb\xbf")
            disp = csv_res.headers.get("content-disposition", "")
            assert "inbox-export-" in disp
            assert ".csv" in disp
            text = csv_res.content.decode("utf-8-sig")
            assert "Id" in text
            assert "ИИ ошибся" in text
            assert "Заметка об ошибке ИИ" in text
            assert visible_id in text
            assert wrong_id in text
            assert hidden_id not in text
            assert "модель ошиблась" in text
            assert "да" in text

            formula = client.post(
                "/api/inbox/export",
                json={
                    "format": "csv",
                    "columns": ["title"],
                    "q": "экспорт visible",
                },
            )
            assert formula.status_code == 200

            # inject formula-like title
            with smoke_db() as session:
                lot = session.get(Lot, visible_id)
                assert lot is not None
                lot.title = "=1+1"
                session.commit()

            formula_res = client.post(
                "/api/inbox/export",
                json={
                    "format": "csv",
                    "columns": ["tender_id", "title"],
                    "q": visible_id,
                },
            )
            assert formula_res.status_code == 200
            ftext = formula_res.content.decode("utf-8-sig")
            assert "'=1+1" in ftext

            xlsx_res = client.post(
                "/api/inbox/export",
                json={
                    "format": "xlsx",
                    "columns": list(AI_REVIEW_PRESET),
                    "ai_wrong": True,
                },
            )
            assert xlsx_res.status_code == 200
            assert xlsx_res.content[:2] == b"PK"
            assert ".xlsx" in (xlsx_res.headers.get("content-disposition") or "")
            from io import BytesIO

            wb = load_workbook(BytesIO(xlsx_res.content), read_only=True)
            ws = wb.active
            assert ws is not None
            rows = list(ws.iter_rows(values_only=True))
            assert rows[0][0] == "Id"
            body_ids = {str(r[0]) for r in rows[1:]}
            assert wrong_id in body_ids
            assert visible_id not in body_ids
            assert hidden_id not in body_ids
            wb.close()
    finally:
        _cleanup(smoke_db, username=username, lot_ids=lot_ids, query=query)
